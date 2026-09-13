"""Generate the app's mock data and Supabase seed from the wedding workbook."""
from __future__ import annotations

import json
import sys
import uuid
from pathlib import Path

import openpyxl


def clean(value):
    return str(value).strip() if value is not None else ""


def sql(value):
    if value is None or value == "":
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


source, project = Path(sys.argv[1]), Path(sys.argv[2])
sheet = openpyxl.load_workbook(source, data_only=True, read_only=True).active
headers = [clean(value) for value in next(sheet.iter_rows(values_only=True))]
required = ["name", "table_number", "seat_number", "group_name", "notes"]
if headers != required:
    raise SystemExit(f"Expected headers {required}; found {headers}")

records = []
for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
    if not any(value not in (None, "") for value in values):
        continue
    row = dict(zip(headers, values))
    name, table_label, seat_text = clean(row["name"]), clean(row["table_number"]).upper(), clean(row["seat_number"])
    if not name or not table_label or not seat_text:
        raise SystemExit(f"Missing required value on row {row_number}")
    records.append({
        "name": name,
        "table_number": table_label,
        "seat_number": int(float(seat_text)),
        "group_name": clean(row["group_name"]) or None,
        "notes": clean(row["notes"]) or None,
    })

labels = sorted({r["table_number"] for r in records}, key=lambda x: (x != "VIP", int(x) if x.isdigit() else 9999, x))
tables = []
for order, label in enumerate(labels, start=1):
    tables.append({
        "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"wedding-table:{label}")),
        "table_number": label,
        "name": "VIP Family" if label == "VIP" else None,
        "sort_order": order,
    })
table_ids = {t["table_number"]: t["id"] for t in tables}

guests = []
for row_number, record in enumerate(records, start=2):
    guests.append({
        "id": str(uuid.uuid5(uuid.NAMESPACE_URL, f"wedding-guest:{row_number}:{record['name']}:{record['table_number']}:{record['seat_number']}")),
        "name": record["name"],
        "table_id": table_ids[record["table_number"]],
        "table_number": record["table_number"],
        "seat_number": record["seat_number"],
        "group_name": record["group_name"],
        "attendance_status": "pending",
        "arrival_time": None,
        "notes": record["notes"],
        "active": True,
    })

typescript = """import type { Guest, WeddingTable } from \"./types\";\n\n"""
typescript += "export const mockTables = " + json.dumps([{k: v for k, v in t.items() if k != "sort_order"} for t in tables], ensure_ascii=False, indent=2) + " satisfies Omit<WeddingTable, \"guests\">[];\n\n"
typescript += "export const mockGuests = " + json.dumps(guests, ensure_ascii=False, indent=2).replace("null", "null").replace("true", "true") + " satisfies Guest[];\n"
(project / "lib" / "mock-data.ts").write_text(typescript, encoding="utf-8")

seed = ["-- Generated from wedding_guest_list.xlsx. Re-running this seed resets imported guests.\n", "begin;", "delete from public.activity_logs;", "delete from public.guests;", "delete from public.tables;\n"]
seed.append("insert into public.tables (id, table_number, name, sort_order) values")
seed.append(",\n".join(f"({sql(t['id'])}::uuid, {sql(t['table_number'])}, {sql(t['name'])}, {t['sort_order']})" for t in tables) + ";\n")
seed.append("insert into public.guests (id, name, table_id, seat_number, group_name, attendance_status, arrival_time, notes, active) values")
seed.append(",\n".join(f"({sql(g['id'])}::uuid, {sql(g['name'])}, {sql(g['table_id'])}::uuid, {g['seat_number']}, {sql(g['group_name'])}, 'pending', null, {sql(g['notes'])}, true)" for g in guests) + ";\n")
seed.extend(["commit;", ""])
(project / "supabase" / "seed.sql").write_text("\n".join(seed), encoding="utf-8")
print(json.dumps({"guests": len(guests), "tables": labels}, ensure_ascii=False))
